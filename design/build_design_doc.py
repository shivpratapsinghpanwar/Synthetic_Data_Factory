"""Build the SDF design document (docx + figures).

Rebuild with:
    uv run --with python-docx --with matplotlib python design/build_design_doc.py

Outputs design/figures/*.png and design/SDF_Design_Document.docx.
Codename policy: studies are referred to only as cond_a, cond_b, ... here.
"""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle

HERE = Path(__file__).resolve().parent
FIGDIR = HERE / "figures"
FIGDIR.mkdir(exist_ok=True)

INK = "#2b3a55"
BLUE = "#dbe7f6"
GREEN = "#ddefdd"
ORANGE = "#fbe8d3"
PURPLE = "#e9e0f4"
GREY = "#eeeeee"
RED = "#f6dcdc"


def box(ax, x, y, w, h, text, fc=BLUE, fs=9.5, weight="normal", ec=INK):
    ax.add_patch(
        FancyBboxPatch(
            (x, y), w, h,
            boxstyle="round,pad=0.03,rounding_size=0.06",
            fc=fc, ec=ec, lw=1.3, mutation_aspect=1.0,
        )
    )
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center",
            fontsize=fs, weight=weight, color="#1c2436", linespacing=1.35)


def arrow(ax, p1, p2, style="-|>", color=INK, lw=1.5, ls="-", rad=0.0):
    ax.add_patch(
        FancyArrowPatch(
            p1, p2, arrowstyle=style, mutation_scale=15,
            color=color, lw=lw, linestyle=ls,
            connectionstyle=f"arc3,rad={rad}", zorder=5,
        )
    )


def canvas(w, h):
    fig, ax = plt.subplots(figsize=(w, h))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")
    return fig, ax


def save(fig, name):
    fig.savefig(FIGDIR / name, dpi=170, bbox_inches="tight",
                facecolor="white")
    plt.close(fig)
    print("wrote", FIGDIR / name)


# ---------------------------------------------------------------- figure 1
def fig1_overview():
    fig, ax = canvas(13, 6.2)

    box(ax, 1, 62, 15, 26,
        "Private study data\ncond_a ≈ 0.7k imgs\ncond_b ≈ 3.3k imgs\n"
        "(+ public benchmark\ntrack: HAM10000)", fc=GREY)
    box(ax, 1, 22, 15, 22,
        "Deterministic splits\ngrouped · stratified\nval / test stay\n100% real", fc=GREY)
    arrow(ax, (8.5, 62), (8.5, 44))

    # generator arms
    box(ax, 22, 76, 20, 13, "A · MDX\nprior-free DiT, 35–80M\nhard-routed experts", fc=GREEN)
    box(ax, 22, 61, 20, 12, "SD1.5-LoRA · 0.86B\npretrained prior", fc=GREEN)
    box(ax, 22, 46, 20, 12, "SDXL-LoRA · 2.6B\npretrained prior", fc=GREEN)
    box(ax, 22, 27, 20, 15,
        "B · CF-Edit\nedits real normals,\nreuses the MDX ckpt\n(no extra training)", fc=PURPLE)
    ax.text(32, 92, "generator committee (C)", fontsize=10.5, ha="center",
            style="italic", color=INK)

    for ay in (82.5, 67, 52, 34.5):
        arrow(ax, (16, 70), (22, ay), rad=0.08)
        arrow(ax, (42, ay), (49, 61), rad=-0.05)

    box(ax, 49, 52, 16, 18,
        "Candidate pool\nper-image manifest:\nsynthetic=true,\narm, source id", fc=ORANGE)
    arrow(ax, (57, 52), (57, 44))
    box(ax, 49, 26, 16, 18,
        "Quality gate\npixel screen ·\nmemorization audit\n(≥0.985 quarantined)", fc=ORANGE)
    arrow(ax, (65, 35), (71, 35))
    box(ax, 71, 26, 15, 18,
        "Select\ndetector-scored,\ntwo-regime\ncomposition", fc=ORANGE)
    arrow(ax, (86, 35), (89, 35))
    box(ax, 89, 26, 10, 18, "Augment\ntrain split\nonly", fc=ORANGE)

    box(ax, 71, 68, 28, 20,
        "Detector arms\nreal_only · oversample-matched\nclassic-aug · mdx-aug · sd-aug\n"
        "cfe-aug · committee-aug", fc=RED, fs=8.5)
    arrow(ax, (94, 44), (91, 68), rad=-0.1)
    box(ax, 71, 3, 28, 15,
        "Evaluate — paired-seed Δ macro-F1\non the curated REAL test split", fc=RED)
    arrow(ax, (85, 68), (85, 18))
    save(fig, "fig1_overview.png")


# ---------------------------------------------------------------- figure 2
def fig2_mdx():
    fig, ax = canvas(13, 6.6)

    box(ax, 1, 42, 12, 14, "image\nres × res\npixels", fc=GREY)
    arrow(ax, (13, 49), (17, 49))
    box(ax, 17, 42, 13, 14, "patchify\npatch = res/16\n→ 256 tokens", fc=GREY)
    arrow(ax, (30, 49), (34, 49))

    # DiT block container
    ax.add_patch(Rectangle((34, 22), 44, 56, fc="#f7f9fc", ec=INK,
                           lw=1.4, ls="--"))
    ax.text(56, 74.5, "DiT block  ×N   (adaLN-zero modulated)",
            fontsize=10.5, ha="center", weight="bold", color=INK)

    box(ax, 36, 56, 9, 12, "LN\n+\nMSA", fs=9)
    arrow(ax, (45, 62), (48, 62))
    box(ax, 48, 56, 9, 12, "residual\n⊕", fs=9, fc=GREY)
    arrow(ax, (57, 62), (60, 62))
    box(ax, 60, 56, 7, 12, "LN", fs=9)
    arrow(ax, (67, 62), (69.5, 62))
    box(ax, 69.5, 56, 7.5, 12, "⊕ out", fs=9, fc=GREY)

    box(ax, 50, 38, 10, 12, "shared\nFFN", fs=9)
    box(ax, 63, 38, 12, 12, "expert FFN\nlow-rank,\nzero-init up", fs=8.5, fc=GREEN)
    arrow(ax, (62, 56), (54, 50), rad=0.15)
    arrow(ax, (65, 56), (67, 50), rad=-0.1)
    arrow(ax, (58, 50), (71.5, 56), rad=-0.2)
    arrow(ax, (72, 50), (74.5, 56), rad=-0.05)

    box(ax, 36, 24, 26, 10,
        "adaLN-zero: scale / shift / gate\n← MLP(t-emb ⊕ class-emb)",
        fs=8, fc=ORANGE)
    arrow(ax, (49, 34), (49, 56), ls=":")
    arrow(ax, (52, 34), (52, 38), ls=":")

    # conditioning inputs
    box(ax, 8, 6, 16, 10, "timestep t\nsinusoidal emb", fc=ORANGE, fs=9)
    box(ax, 27, 6, 20, 10, "class / condition label\n(vocab lookup, CFG dropout)",
        fc=ORANGE, fs=9)
    arrow(ax, (16, 16), (40, 26), rad=0.15)
    arrow(ax, (37, 16), (44, 26), rad=0.05)

    # study routing + onboarding
    box(ax, 82, 52, 17, 22,
        "expert table\none low-rank\nexpert per study,\nevery block", fc=GREEN, fs=9)
    box(ax, 82, 24, 17, 18, "label vocab\nrows per\nstudy class", fc=ORANGE, fs=9)
    arrow(ax, (82, 58), (75.5, 45), rad=0.15)
    arrow(ax, (82, 30), (47, 12), rad=-0.2)
    box(ax, 55, 4, 44, 12,
        "HARD routing by study id — no learned router.\n"
        "Onboarding: append vocab + expert rows, freeze trunk →\n"
        "previously learned studies stay bitwise unchanged.",
        fs=9, fc=PURPLE)
    arrow(ax, (72, 16), (68, 37), rad=0.1)
    save(fig, "fig2_mdx.png")


# ---------------------------------------------------------------- figure 3
def fig3_cfedit():
    fig, ax = canvas(13, 5.6)

    box(ax, 2, 58, 20, 24,
        "real NORMAL image $x_0$\n(train split only —\nnever val/test)", fc=GREY)
    arrow(ax, (22, 70), (32, 70))
    ax.text(27, 88, "forward diffusion\nto $t = s\\cdot T$", fontsize=9,
            ha="center", color=INK)
    box(ax, 32, 58, 22, 24,
        "noised $x_t$\nglobal structure survives,\nappearance detail\ndestroyed", fc=ORANGE)
    arrow(ax, (54, 70), (66, 70))
    ax.text(60, 91, "reverse denoising with\nTARGET label + CFG,\nrouted through the study's expert",
            fontsize=9, ha="center", color=INK)
    box(ax, 66, 58, 22, 24,
        "counterfactual $\\hat{x}_0$\nsame scene, now\nexhibiting the condition", fc=PURPLE)

    # strength dial
    ax.add_patch(Rectangle((10, 34), 70, 6, fc="#f0f0f0", ec=INK, lw=1.2))
    for s, lab in ((0.3, "0.3"), (0.5, "0.5"), (0.7, "0.7")):
        x = 10 + 70 * s
        ax.plot([x, x], [34, 40], color=INK, lw=1.4)
        ax.text(x, 30.5, lab, ha="center", fontsize=9, color=INK)
    ax.text(10, 28, "low strength s:\nkeeps the real image —\nrealism ↑, subtle anomaly",
            fontsize=8.5, ha="left", va="top", color=INK)
    ax.text(80, 28, "high strength s:\napproaches pure generation",
            fontsize=8.5, ha="right", va="top", color=INK)
    ax.text(45, 44.5, "strength dial  s ∈ (0,1)", fontsize=10, ha="center",
            weight="bold", color=INK)

    box(ax, 25, 2, 50, 13,
        "output: minimal-difference pair  ($x_0$, $\\hat{x}_0$)\n"
        "straddles the decision boundary; source id recorded in the manifest",
        fc=RED, fs=9.5)
    arrow(ax, (77, 58), (60, 15), rad=-0.15)
    arrow(ax, (12, 58), (35, 15), rad=0.15)
    save(fig, "fig3_cfedit.png")


# ---------------------------------------------------------------- figure 4
def fig4_select():
    fig, ax = canvas(11, 5.0)

    for y0, title, keep, keeplab, extra in (
        (58, "pure-generation arms (MDX / SD-LoRA samples)",
         (72, 98), "KEEP: top confidence\n= label correctness", None),
        (18, "CF-Edit arm (counterfactual pairs)",
         (48, 84), "KEEP: boundary band,\nrequire  p(target) − p(source) ≥ δ",
         None),
    ):
        ax.text(4, y0 + 26, title, fontsize=10.5, weight="bold", color=INK)
        ax.add_patch(Rectangle((4, y0), 94, 8, fc="#f4f4f4", ec=INK, lw=1.1))
        ax.add_patch(Rectangle((keep[0], y0), keep[1] - keep[0], 8,
                               fc=GREEN, ec=INK, lw=1.1))
        ax.add_patch(Rectangle((4, y0), 22, 8, fc=RED, ec=INK, lw=1.1))
        ax.text(15, y0 + 4, "discard:\nwrong / weak label", fontsize=7.5,
                ha="center", va="center")
        ax.text((keep[0] + keep[1]) / 2, y0 + 12.5, keeplab, fontsize=8.5,
                ha="center", color="#1e5c1e")
        for x, lab in ((4, "0"), (51, "0.5"), (98, "1")):
            ax.text(x, y0 - 4, lab, fontsize=8, ha="center", color=INK)
        if extra:
            ax.text(51, y0 - 9.5, extra, fontsize=8, ha="center",
                    style="italic", color="#555555")
    ax.text(51, 2, "x-axis: real-only detector score  p(target class | image)",
            fontsize=9.5, ha="center", color=INK)
    save(fig, "fig4_select.png")


# ---------------------------------------------------------------- figure 5
def fig5_roadmap():
    fig, ax = canvas(13, 4.6)

    ms = [
        ("M2b\nSD1.5 arm on cond_a\n(collect pending)", ORANGE),
        ("M3\nMDX samples +\nCF-Edit strength sweep\n+ quality gates", BLUE),
        ("M4\ndetector arms:\nreal_only, controls,\nper-arm augmented", BLUE),
        ("M5\ncommittee select +\nmulti-seed replication\n(3–5 seeds)", BLUE),
        ("M6\n256px winner rerun +\nexpert-placement /\nband ablations", BLUE),
        ("Paper freeze\nfigures, tables,\nrepro package", GREEN),
    ]
    x = 1
    for text, fc in ms:
        box(ax, x, 52, 14.5, 30, text, fc=fc, fs=8.5)
        if x > 1:
            arrow(ax, (x - 2, 67), (x, 67))
        x += 16.5

    box(ax, 10, 12, 38, 22,
        "public benchmark lane (runs in parallel):\nHAM10000 rare classes + one public\n"
        "paired normal/abnormal dataset → the\nreproducible results reviewers can rerun",
        fc=PURPLE, fs=9)
    arrow(ax, (48, 23), (86, 52), rad=-0.15)
    box(ax, 62, 12, 30, 16,
        "compute budget: single T4,\n30 h GPU / week — matrix is\npruned, seeds before scale",
        fc=GREY, fs=9)
    save(fig, "fig5_roadmap.png")


# ---------------------------------------------------------------- document
def build_doc():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)

    def h(text, level=1):
        doc.add_heading(text, level=level)

    def p(text, bold=False, italic=False):
        par = doc.add_paragraph()
        run = par.add_run(text)
        run.bold = bold
        run.italic = italic
        return par

    def bullets(items, style="List Bullet"):
        for it in items:
            doc.add_paragraph(it, style=style)

    def fig(name, caption, width=6.4):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(str(FIGDIR / name), width=Inches(width))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        for j, htxt in enumerate(headers):
            cell = t.rows[0].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(htxt)
            run.bold = True
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                c = t.rows[i].cells[j]
                c.text = str(val)
                for par in c.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(9.5)
        if widths:
            for j, wd in enumerate(widths):
                for row in t.rows:
                    row.cells[j].width = Inches(wd)
        doc.add_paragraph()

    # ---- title page ----
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tp.add_run("\nSynthetic Data Factory\n")
    run.bold = True
    run.font.size = Pt(26)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "Design document — MDX, CF-Edit and the generator committee\n"
        "A research program toward a tier-1 publication\n")
    run.font.size = Pt(13)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Version 1.0 · 2026-09-06 · studies referenced by "
                       "codename (cond_a, cond_b, …) throughout")
    run.italic = True
    run.font.size = Pt(10)

    doc.add_paragraph()
    h("Contents", level=2)
    bullets([
        "1. Executive summary",
        "2. Problem setting and constraints",
        "3. System overview",
        "4. Architecture A — MDX: mixture-of-experts diffusion transformer",
        "5. Architecture B — CF-Edit: counterfactual anomaly synthesis",
        "6. Architecture C — generator committee and two-regime selection",
        "7. Design resolutions (this iteration)",
        "8. Evaluation protocol",
        "9. Experimental matrix and compute budget",
        "10. Path to a tier-1 paper",
        "11. Risks and mitigations",
        "12. Current empirical state and roadmap",
    ], style="List Number")
    doc.add_page_break()

    # ---- 1 ----
    h("1. Executive summary")
    p("The Synthetic Data Factory turns small private clinical imaging "
      "studies (hundreds to a few thousand images each) into measurable "
      "detector improvements by generating, auditing and selecting "
      "synthetic training images. The design has three layers:")
    bullets([
        "MDX (A): one from-scratch pixel-space diffusion transformer trained "
        "jointly across studies, with hard-routed per-study low-rank experts. "
        "It gives cross-study transfer under scarcity and a frozen-trunk "
        "onboarding guarantee — new studies are added without touching, even "
        "bitwise, what was already learned.",
        "CF-Edit (B): the paper's primary claim. Instead of generating "
        "anomalous images from pure noise, partially noise a REAL normal "
        "image and denoise it under the anomalous label. The result is a "
        "minimal-difference counterfactual pair that straddles the decision "
        "boundary — precisely the examples detectors learn most from — at "
        "zero additional training cost.",
        "Committee + select (C): heterogeneous generator arms across the "
        "prior spectrum (prior-free 35–80M MDX vs pretrained 0.86B / 2.6B "
        "LoRA arms), pooled and composed by a detector-scored selection "
        "stage. This converts 'which prior should synthetic medical data "
        "come from?' from a bet into an ablation table.",
    ])
    p("Primary publishable claim: counterfactual pairs from label-swapped "
      "diffusion editing beat pure generation for rare-condition detector "
      "training under extreme data scarcity. Secondary: the frozen-trunk "
      "onboarding guarantee as a systems contribution. Target venues: "
      "CVPR / ICML / NeurIPS main track, with MICCAI/MIDL as the "
      "domain-venue fallback; a public-dataset reproducibility lane "
      "(HAM10000 rare classes + one more public paired dataset) makes the "
      "results verifiable by reviewers.", bold=False)

    # ---- 2 ----
    h("2. Problem setting and constraints")
    bullets([
        "Data: private multi-study clinical datasets, roughly 0.3k–3.3k "
        "images per study, with curated evaluation splits. Studies are "
        "referred to by codename (cond_a, cond_b, …). Real class names, "
        "dataset slugs and paths live only in gitignored local overlays.",
        "Hard metric: paired-seed detector lift — macro-F1 on the curated, "
        "100% real test split, real-only vs augmented, identical seeds.",
        "Compute: a single 16 GB T4 GPU session at a time, ~30 h GPU/week "
        "shared with other projects. Every experiment must justify its "
        "GPU minutes.",
        "Privacy: client data lives only in private Kaggle datasets "
        "end-to-end; only pipeline code is public. The memorization audit "
        "is a mandatory OUTPUT gate (cosine similarity ≥ 0.985 to any "
        "training image ⇒ quarantined by the augment stage) — required "
        "both for privacy and for any augmentation claim to be "
        "scientifically valid.",
        "Splits: grouped (no near-duplicate leakage across splits), "
        "stratified, deterministic. Synthetic data only ever augments the "
        "train split; val and test stay real forever.",
    ])

    # ---- 3 ----
    h("3. System overview")
    p("The pipeline is staged (audit → train_joint / train_lora → sample / "
      "counterfact → quality_gate → select → augment → train_detector → "
      "evaluate), each stage a resumable CLI entrypoint executed on Kaggle "
      "by the kernel runner against the pushed git revision. Every "
      "synthetic image carries a manifest row (synthetic=true, generator "
      "arm, and — for CF-Edit — the source image id).")
    fig("fig1_overview.png",
        "Figure 1 — End-to-end system. Generator arms feed one audited pool; "
        "selection composes the augmentation set; evaluation is always "
        "paired-seed macro-F1 on real test data.")

    # ---- 4 ----
    h("4. Architecture A — MDX: mixture-of-experts diffusion transformer")
    p("MDX is a from-scratch pixel-space DiT (~35M parameters at dim 384; "
      "~80M at dim 512 / depth 16) trained jointly on every study's train "
      "split. Patch size is chosen as resolution/16, so the sequence is "
      "always 256 tokens at any resolution — the same trunk trains at "
      "64/128/256 px without architectural change.")
    fig("fig2_mdx.png",
        "Figure 2 — One MDX block. The shared FFN carries general image "
        "statistics; a per-study low-rank expert (zero-initialized up-"
        "projection) carries study-specific signal and is hard-routed by "
        "study id. adaLN-zero modulates every block from (timestep ⊕ class) "
        "embeddings.")
    p("Design rationale:", bold=True)
    bullets([
        "Joint trunk — cross-study transfer under scarcity. A 0.7k-image "
        "study borrows general image statistics (edges, anatomy-scale "
        "structure, lighting) from a 3.3k-image one. Tested against an "
        "equal-compute per-class DDPM baseline.",
        "Hard routing, not learned soft MoE routing. Tiny datasets cannot "
        "estimate a router, and a mis-trained router silently corrupts "
        "every study. The study id is known at train and sample time, so "
        "routing needs no learning at all.",
        "Zero-initialized expert up-projections. At step 0 each expert is "
        "an exact no-op, so early training is pure trunk learning; experts "
        "grow into the residual only as study-specific signal warrants.",
        "Frozen-trunk onboarding — a hard guarantee. New studies append "
        "rows to the label vocab and expert tables; with freeze_trunk=1 "
        "only the new rows receive gradients, so previously learned "
        "studies are BITWISE unchanged. Most continual-learning work can "
        "only approximate this. Onboarding cost: minutes.",
        "Classifier-free guidance via label dropout (the expert route is "
        "kept during dropout — the null token is per-study, so guidance "
        "sharpens class identity, not study identity).",
        "EMA weights for sampling; self-describing checkpoints (vocab + "
        "architecture JSON embedded) so a checkpoint alone is sufficient "
        "to resume, onboard, or sample.",
    ])
    p("Positioning: eDiff-I ensembles denoisers over timestep bands; "
      "DiT-MoE scales token-level soft-routed experts for capacity. "
      "Neither targets the multi-dataset scarce-data regime, and neither "
      "gives the frozen-trunk onboarding guarantee. The closest neighbor "
      "is adapter-per-task continual diffusion; our delta is shared-trunk "
      "joint pretraining across private clinical studies plus study-level "
      "expert routing.", italic=True)

    # ---- 5 ----
    h("5. Architecture B — CF-Edit: counterfactual anomaly synthesis")
    p("The highest-leverage observation in this problem: clinical datasets "
      "come with paired class structure — the same anatomy imaged the same "
      "way, with and without the condition. Generating an anomalous image "
      "from pure noise wastes model capacity reproducing what real data "
      "already provides (pose, lighting, device signature, background). "
      "CF-Edit instead:")
    bullets([
        "1. takes a REAL image of the normal class (train split only);",
        "2. diffuses it forward to an intermediate timestep t = s·T "
        "(SDEdit-style partial noising: global structure survives, "
        "appearance detail is destroyed);",
        "3. denoises with the TARGET (anomalous) label under CFG, routed "
        "through the study's expert.",
    ], style="List Number")
    fig("fig3_cfedit.png",
        "Figure 3 — CF-Edit. The strength dial s interpolates between "
        "'barely edited real image' and 'pure generation'; each output is "
        "paired with its source in the manifest.")
    p("Why this should beat pure generation for detector training:", bold=True)
    bullets([
        "Sample efficiency — the generator only has to know the DIFFERENCE "
        "between classes, exactly what the joint conditional model "
        "concentrates in its class embeddings and experts.",
        "Hard pairs by construction — each real normal yields a minimal-"
        "difference (normal, anomalous) pair straddling the decision "
        "boundary; the recorded pairing also enables paired training "
        "objectives later (contrastive or margin losses over pairs).",
        "Label correctness is auditable — the select stage scores every "
        "counterfactual with the real-only detector; edits too weak to "
        "express the condition are filtered before training, and the "
        "flagged rate itself measures the generator's class separation.",
        "Cost — zero additional training (reuses the MDX checkpoint) and "
        "LOWER sampling cost than pure generation (only the last "
        "s-fraction of the reverse trajectory is run).",
    ])
    p("Positioning: SDEdit edits with a prompt, not a learned per-study "
      "class swap; counterfactual-explanation methods produce per-"
      "classifier gradients, not training data; AnoDDPM-style medical work "
      "runs the opposite direction (remove the anomaly to detect it). "
      "Using label-swapped partial denoising of real normals as a data "
      "factory for rare-condition detectors, on sub-1k private datasets, "
      "with a memorization audit, is an open slot in the literature.",
      italic=True)

    # ---- 6 ----
    h("6. Architecture C — generator committee and two-regime selection")
    p("Heterogeneous arms span the prior spectrum: prior-free MDX (no "
      "pretraining leakage, data-hungry) vs pretrained LoRA arms at ~0.86B "
      "(SD1.5) and ~2.6B (SDXL) (sample-efficient, but import a natural-"
      "image prior of unknown relevance to clinical pixels). All arms "
      "sample into one pool; the select stage composes the augmentation "
      "set. CF-Edit slots in as just another arm.")
    p("Two-regime selection (design resolution R1 — see §7): a single "
      "best-first ranking by detector confidence would systematically "
      "discard exactly the boundary-straddling images CF-Edit exists to "
      "produce. Selection therefore runs per-regime:", bold=False)
    fig("fig4_select.png",
        "Figure 4 — Selection regimes. Pure-generation arms keep top-"
        "confidence samples (confidence = label correctness). CF-Edit "
        "keeps a boundary band, requiring the edit to have moved class "
        "by a margin δ while staying informative.", width=5.6)
    bullets([
        "Pure-generation arms: rank by p(target); keep best-first. High "
        "confidence means the label is right — the sample teaches the "
        "class prototype.",
        "CF-Edit arm: keep if p(target) − p(source) ≥ δ (the edit "
        "expressed the condition), then PREFER samples closest to the "
        "decision boundary (|p(target) − 0.5| small). These are the hard "
        "examples; trivially easy edits add little.",
        "Round-robin blending across arms best-first within each arm's "
        "regime, up to the per-class augmentation budget.",
    ])
    p("The empirical question the committee answers — where along the "
      "prior spectrum does useful synthetic medical data come from at "
      "each dataset size? — is unresolved and practically important: "
      "privacy and licensing arguments favor prior-free generation, "
      "quality arguments favor foundation priors. The committee makes it "
      "an ablation table instead of a bet.")

    # ---- 7 ----
    h("7. Design resolutions (this iteration)")
    table(
        ["#", "Tension", "Resolution"],
        [
            ["R1",
             "Best-first selection contradicts CF-Edit's value "
             "proposition (it would drop boundary-straddling pairs).",
             "Two-regime select (§6): confidence-ranked for pure "
             "generation; margin-then-boundary-band for CF-Edit. "
             "Implement before M3 burns GPU."],
            ["R2",
             "Experts are FFN-only. For structural conditions the class "
             "difference may live in early (structure-stage) denoising "
             "that FFN experts under-serve.",
             "Cheap probe first: if CF-Edit at s=0.3 cannot move class "
             "on cond_a, that is evidence of a structure-stage capacity "
             "bottleneck → then ablate attention-path experts and "
             "timestep-band-sliced experts (architecture D) at equal "
             "rank budget."],
            ["R3",
             "Committee comparisons at mixed resolutions (MDX 128px vs "
             "SD 512px) conflate architecture with resolution.",
             "Run the CF-Edit strength sweep at 128px first (cheap); "
             "replicate the winning configuration at 256px before any "
             "cross-arm table is reported. All detector arms train at "
             "one fixed input resolution regardless of generator "
             "output size."],
            ["R4",
             "A counterfactual is legitimately similar to its own "
             "source image — a naive memorization audit would flag "
             "every CF-Edit output.",
             "The audit excludes the recorded source image for CF-Edit "
             "outputs, but ALSO audits against the source class: an "
             "edit that failed to move class is a near-copy by design "
             "and must be caught. Verify the counterfact manifest/gate "
             "wiring before M3."],
        ],
        widths=[0.4, 2.6, 3.7],
    )

    # ---- 8 ----
    h("8. Evaluation protocol")
    bullets([
        "Paired-seed detector deltas: real-only vs each augmented arm, "
        "identical seeds, reported as mean ± sd of macro-F1 on the "
        "curated real test split, with a paired test across seeds "
        "(3 seeds minimum for screening, 5 for reported tables).",
        "Balance-matched real-oversampling control arm — so lift cannot "
        "be explained by class-balance shift alone (this control caught "
        "an off-target melanoma-recall artifact in the HAM10000 pilot).",
        "Classical-augmentation control arm (standard photometric + "
        "geometric policy) — so lift cannot be explained by generic "
        "augmentation effects.",
        "Memorization audit on every reported set: max cosine similarity "
        "to train embeddings, flagged ≥ 0.985 quarantined; for CF-Edit "
        "the audit runs with source-exclusion + source-class check (R4). "
        "The flagged rate is itself reported.",
        "Per-class F1 deltas reported alongside macro-F1 — rare-class "
        "lift is the actual product; macro movement alone can hide "
        "off-target effects.",
    ])

    # ---- 9 ----
    h("9. Experimental matrix and compute budget")
    p("Detector arms × datasets. Every cell is a paired-seed Δ macro-F1 "
      "against the same real_only baseline; controls isolate the "
      "explanation.")
    table(
        ["Arm", "What it isolates", "cond_a", "HAM-rare (public)", "cond_b"],
        [
            ["real_only", "baseline", "0.9216 (done)", "done (pilot)",
             "0.9897 — at ceiling, excluded"],
            ["oversample-matched", "class-balance shift", "planned",
             "planned", "—"],
            ["classic-aug", "generic augmentation", "planned", "planned", "—"],
            ["mdx-aug", "prior-free generation", "planned", "planned", "—"],
            ["sd15-aug", "0.86B pretrained prior", "M2b in flight",
             "done (pilot: vasc F1 +0.050±0.013)", "—"],
            ["sdxl-aug", "2.6B pretrained prior", "stretch", "stretch", "—"],
            ["cfe-aug (s sweep)", "counterfactual pairs vs pure gen",
             "planned (primary claim)", "planned", "—"],
            ["committee-aug", "selection across arms", "planned",
             "planned", "—"],
        ],
        widths=[1.1, 1.7, 1.4, 1.5, 1.0],
    )
    p("cond_b's real-only baseline is already at ceiling (0.9897 "
      "macro-F1), so it contributes to MDX joint-training transfer "
      "(bigger sibling study) and to the onboarding demonstration, not "
      "to augmentation claims. Compute discipline: strength sweeps and "
      "screening at 128px (~20 min/run on T4); only winners replicate "
      "at 256px and across 5 seeds. Estimated total for the matrix: "
      "roughly 25–35 T4-hours, spread over weekly budget.")

    # ---- 10 ----
    h("10. Path to a tier-1 paper")
    p("What a CVPR/ICML/NeurIPS reviewer will demand, and how the design "
      "answers it:", bold=True)
    table(
        ["Reviewer demand", "Our answer"],
        [
            ["A crisp, falsifiable novel claim",
             "CF-Edit: counterfactual pairs beat pure generation for "
             "rare-condition detector training under scarcity — tested "
             "arm-vs-arm with paired seeds on the same pool budget."],
            ["Reproducibility on public data",
             "A public lane run with the identical pipeline: HAM10000 "
             "rare classes (akiec/df/vasc) + one additional public "
             "paired normal/abnormal dataset (candidate: a chest-X-ray "
             "normal/pneumonia set — public, paired structure, license-"
             "friendly). Private studies become the in-the-wild "
             "evidence; public lane is what reviewers rerun."],
            ["Strong baselines",
             "Not just real_only: balance-matched oversampling, classical "
             "augmentation policy, pretrained-LoRA generation at two "
             "scales, prior-free generation, and (stretch) a GAN "
             "baseline (StyleGAN2-ADA is the standard small-data GAN)."],
            ["Ablations",
             "Strength sweep (s = 0.3/0.5/0.7), selection policy "
             "(best-first vs two-regime), expert placement (FFN vs "
             "attention vs timestep-band, equal rank budget), joint vs "
             "single-study trunk, resolution (128 vs 256), synthetic "
             "volume dose–response."],
            ["Statistical rigor",
             "Paired seeds (5 for reported tables), mean ± sd, paired "
             "test; per-class deltas; controls that kill the two "
             "obvious alternative explanations."],
            ["Privacy story (medical data)",
             "Private data never leaves private storage; memorization "
             "audit with quarantine is part of the pipeline, and the "
             "frozen-trunk guarantee bounds cross-study contamination "
             "at onboarding time."],
        ],
        widths=[1.9, 4.8],
    )
    p("Venue plan (today = 2026-09-06):", bold=True)
    table(
        ["Venue", "Deadline (approx.)", "Fit / call"],
        [
            ["CVPR 2027", "~mid-Nov 2026",
             "Aggressive: ~10 weeks for the full matrix + writing. "
             "Feasible only if M3–M5 land clean on the first pass. "
             "Treat as a stretch target."],
            ["ICML 2027", "~late Jan 2027",
             "Realistic primary target: matrix + 5-seed replication + "
             "public lane fit comfortably."],
            ["NeurIPS 2027", "~May 2027",
             "Comfortable fallback; also fits the Datasets & Benchmarks "
             "track if the factory framing outgrows the method framing."],
            ["MICCAI 2027", "~Feb–Mar 2027",
             "Domain fallback with high acceptance odds for this topic; "
             "a MIDL/MICCAI-workshop version can precede the main-"
             "conference push without burning novelty."],
        ],
        widths=[1.0, 1.3, 4.4],
    )
    bullets([
        "Anonymity note: the public repo already describes the research "
        "program. For double-blind submission either make the repo "
        "private for the review window or scrub the paper of links; "
        "decide before the first deadline.",
        "Framing note: lead with CF-Edit (the claim), present MDX as the "
        "enabling model and the committee as the measurement instrument. "
        "The frozen-trunk guarantee is one strong section, not the title.",
    ])

    # ---- 11 ----
    h("11. Risks and mitigations")
    table(
        ["Risk", "Signal", "Mitigation"],
        [
            ["CF-Edit cannot move class at low strength",
             "s=0.3 probe: p(target) ≈ p(source)",
             "Raise s; R2 expert-placement ablation; more joint-training "
             "steps; fall back to claiming the committee/selection result."],
            ["Memorization on sub-1k studies",
             "audit flagged-rate high",
             "Quarantine is automatic; reduce steps / add dropout; report "
             "flagged rate honestly — it strengthens the privacy story."],
            ["Detector ceiling too close (cond_a 0.9216)",
             "deltas within seed noise",
             "Per-class rare-condition F1 is the headline metric, not "
             "macro; scarcity ablation (train on 25%/50% subsets) "
             "widens headroom and IS the paper's regime anyway."],
            ["Compute contention (shared 30 h/week)",
             "GPU slots held by other kernels",
             "Screen at 128px; never two concurrent kernel versions; "
             "queue via `kernels status` checks; Colab/cloud adapter is "
             "the approved overflow route."],
            ["Pretrained arms leak natural-image bias into medical pixels",
             "sd/sdxl samples pass gate but hurt detector",
             "That result is itself a finding the committee is designed "
             "to surface; selection filters it operationally."],
            ["Scooping (public repo describes the ideas)",
             "—",
             "Workshop paper early to timestamp the idea, or repo "
             "private until submission (owner's call, see §10)."],
        ],
        widths=[1.9, 1.6, 3.2],
    )

    # ---- 12 ----
    h("12. Current empirical state and roadmap")
    table(
        ["Milestone", "Status", "Key numbers"],
        [
            ["HAM10000 pilot (M5, public lane)", "done",
             "3 seeds × 2 arms, 400 gated synthetic images; vasc F1 "
             "+0.050 ± 0.013; balance-shift artifact caught by control."],
            ["Real-only baselines (client lane)", "done",
             "cond_a macro-F1 0.9216 (the target with headroom); "
             "cond_b 0.9897 (ceiling — excluded from augmentation)."],
            ["MDX CPU test suite", "done", "12/12 on Kaggle CPU."],
            ["M1 joint sanity (64px)", "done", "loss 0.442 → 0.065, 33.8M."],
            ["M2 joint trunk (128px)", "done",
             "79M trunk, 6000 steps, loss 0.489 → 0.0446, 22 min T4; "
             "64/64 samples pass pixel screen; checkpoint published to "
             "the artifacts dataset."],
            ["M2b SD1.5 heavy arm on cond_a", "in flight",
             "1200 steps @ 512px, batch 2 × accum 2 (batch 16 OOMs a "
             "T4 — proven); collect pending."],
            ["M3 samples + CF-Edit sweep", "next",
             "~100/class MDX samples + counterfactuals at s = "
             "0.3/0.5/0.7 from the published checkpoint + per-class "
             "quality gates. Blocked on R1 + R4 code first."],
        ],
        widths=[1.9, 0.8, 4.0],
    )
    fig("fig5_roadmap.png",
        "Figure 5 — Roadmap. The public benchmark lane runs the identical "
        "pipeline and produces the reviewer-reproducible results; private "
        "studies are the in-the-wild evidence.")
    p("GPU spent this cycle: ~1 h of the 30 h/week budget. The design's "
      "guiding discipline stands: screen cheap, replicate winners, and "
      "never report a number without its control.")

    out = HERE / "SDF_Design_Document.docx"
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    fig1_overview()
    fig2_mdx()
    fig3_cfedit()
    fig4_select()
    fig5_roadmap()
    build_doc()
